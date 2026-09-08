"""Export slides/index.html to a 16:9 PowerPoint deck.

The HTML deck is the source of truth. This script renders each slide with
headless Chrome at 2x and places the images full-bleed on 13.333 x 7.5 in
slides, so the .pptx is pixel-identical to what the browser shows. Speaker
notes are carried over as real PowerPoint notes.

The result is an image deck: faithful, but not text-editable in PowerPoint.
Edit `slides/index.html` and re-export. It writes `slides/deck.pptx` and never
touches `Thanh_Vo_KYC_Agent_Demo.pptx`; pass `--out` to choose another path.

Requirements: Google Chrome (or set CHROME=/path/to/chrome) and python-pptx.

    python3 slides/to_pptx.py             # -> slides/deck.pptx
    python3 slides/to_pptx.py --keep-png  # keep slides/.export/*.png
"""
from __future__ import annotations

import argparse
import html
import re
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
DECK = HERE / "index.html"
EXPORT_DIR = HERE / ".export"
DEFAULT_OUT = HERE / "deck.pptx"

WIDTH_IN, HEIGHT_IN = 13.3333, 7.5
SCALE = 2  # 2560x1440 per slide

CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/Applications/Brave Browser.app/Contents/MacOS/Brave Browser",
    "google-chrome",
    "chromium",
]


def find_chrome() -> str:
    import os

    env = os.environ.get("CHROME")
    if env:
        return env
    for candidate in CHROME_CANDIDATES:
        if Path(candidate).exists() or shutil.which(candidate):
            return candidate
    sys.exit("no Chrome found; set CHROME=/path/to/chrome")


def parse_deck() -> list[dict[str, str]]:
    """Slide count plus speaker notes, read straight out of the HTML."""
    src = DECK.read_text()
    slides = re.findall(r'<section class="slide">(.*?)</section>', src, re.S)
    if not slides:
        sys.exit(f"no slides found in {DECK}")
    out = []
    for markup in slides:
        note = re.search(r'<aside class="notes"([^>]*)>(.*?)</aside>', markup, re.S)
        text, sources = "", ""
        if note:
            attrs, inner = note.groups()
            text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", inner)).strip()
            src_attr = re.search(r'data-src="([^"]*)"', attrs)
            sources = html.unescape(src_attr.group(1)) if src_attr else ""
        out.append({"notes": html.unescape(text), "sources": sources})
    return out


def png_size(path: Path) -> tuple[int, int]:
    """Width/height straight from the IHDR chunk -- no image library needed."""
    head = path.read_bytes()[:24]
    if head[:8] != b"\x89PNG\r\n\x1a\n" or head[12:16] != b"IHDR":
        raise ValueError(f"{path} is not a PNG")
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")


def shoot(chrome: str, index: int, target: Path) -> None:
    target.unlink(missing_ok=True)  # never let a stale render survive a failure
    url = f"{DECK.as_uri()}?capture#{index}"
    proc = subprocess.run(
        [
            chrome,
            "--headless",
            "--disable-gpu",
            "--hide-scrollbars",
            "--no-first-run",
            "--virtual-time-budget=2500",
            f"--force-device-scale-factor={SCALE}",
            "--window-size=1280,720",
            f"--screenshot={target}",
            url,
        ],
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0 or not target.exists():
        sys.exit(f"chrome failed on slide {index} (exit {proc.returncode}):\n{proc.stderr[-800:]}")
    want = (1280 * SCALE, 720 * SCALE)
    got = png_size(target)
    if got != want:
        sys.exit(f"slide {index} rendered at {got}, expected {want}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--keep-png", action="store_true", help="keep the rendered PNGs")
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT, help="output .pptx path")
    args = ap.parse_args()

    chrome = find_chrome()
    slides = parse_deck()
    shutil.rmtree(EXPORT_DIR, ignore_errors=True)
    EXPORT_DIR.mkdir()

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH_IN), Inches(HEIGHT_IN)
    blank = prs.slide_layouts[6]

    for i, meta in enumerate(slides, start=1):
        png = EXPORT_DIR / f"slide-{i:02d}.png"
        shoot(chrome, i, png)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, Inches(WIDTH_IN), Inches(HEIGHT_IN))
        if meta["notes"]:
            body = meta["notes"]
            if meta["sources"]:
                body += f"\n\nSources: {meta['sources']}"
            slide.notes_slide.notes_text_frame.text = body
        print(f"slide {i:02d}/{len(slides)}  {png.name}")

    prs.save(args.out)
    if not args.keep_png:
        shutil.rmtree(EXPORT_DIR, ignore_errors=True)
    print(f"wrote {args.out} — {len(slides)} slides at {WIDTH_IN}x{HEIGHT_IN} in")


if __name__ == "__main__":
    main()
